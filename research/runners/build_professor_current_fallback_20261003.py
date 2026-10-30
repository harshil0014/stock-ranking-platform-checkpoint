"""Create a current-only, fill-missing professor workbook overlay.

The accepted company-fact snapshot and the source workbook are read-only.
No historical panel, baseline score, or production pointer is modified.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import csv
import gzip
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(r"C:\Users\Harshil\Downloads\_Organized\Documents\Copy of OK Strategy - Master Sheet v4.3 (1).xlsx")
ACCEPTED_ROOT = ROOT / "data/company_facts"
OUT = ROOT / "experiments/professor-current-fallback-20261003"
FIELD_SOURCES = {
    "professor_workbook_cagr_10y": ("CAGR", "D"),
    "professor_workbook_cagr_5y": ("CAGR", "E"),
    "professor_workbook_cagr_3y": ("CAGR", "F"),
    "professor_workbook_cagr_1y": ("CAGR", "G"),
    "professor_workbook_current_price": ("CAGR", "C"),
    "professor_workbook_market_cap": ("CAGR", "I"),
    "professor_industry": ("CAGR", "H"),
    "professor_workbook_ath_price": ("ATH", "D"),
    "professor_workbook_ath_date": ("ATH", "E"),
    "professor_workbook_cagr_7y": ("ATH_RETURNS", "L"),
}
NUMERIC = set(FIELD_SOURCES) - {"professor_industry", "professor_workbook_ath_date"}
RETURN_FIELDS = {"professor_workbook_cagr_10y": "K", "professor_workbook_cagr_7y": "L",
                 "professor_workbook_cagr_5y": "M", "professor_workbook_cagr_3y": "N",
                 "professor_workbook_cagr_1y": "O", "professor_industry": "J"}
STOP = {"LTD", "LIMITED", "INDIA", "THE", "PRIVATE", "PVT", "CO", "COMPANY"}


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def blank(value: object) -> bool:
    if value is None:
        return True
    try:
        if bool(pd.isna(value)):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def norm_name(value: object) -> str:
    words = re.findall(r"[A-Z0-9]+", str(value or "").upper().replace("&", " AND "))
    return " ".join(w for w in words if w not in STOP)


def name_score(a: object, b: object) -> float:
    aa, bb = norm_name(a), norm_name(b)
    if not aa or not bb:
        return 0.0
    if aa == bb:
        return 1.0
    ratio = SequenceMatcher(None, aa, bb).ratio()
    wa, wb = set(aa.split()), set(bb.split())
    overlap = len(wa & wb) / min(len(wa), len(wb))
    return ratio if overlap >= 0.5 else 0.0


def clean_code(value: object) -> str:
    text = str(value or "").strip().upper()
    return text.replace("BSE:", "BOM:") if text else ""


def valid(field: str, value: object, max_date: date) -> object | None:
    if blank(value) or isinstance(value, str) and value.startswith("#"):
        return None
    if field == "professor_industry":
        return str(value).strip() if str(value).strip() else None
    if field == "professor_workbook_ath_date":
        try:
            day = pd.Timestamp(value).date()
        except (TypeError, ValueError, OverflowError):
            return None
        return day.isoformat() if date(1990, 1, 1) <= day <= max_date else None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    if "cagr" in field and number < -1:
        return None
    if ("price" in field or "market_cap" in field) and number <= 0:
        return None
    return number


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    state_file = ACCEPTED_ROOT / "accepted_state.json"
    state = json.loads(state_file.read_text(encoding="utf-8"))
    source_csv = ACCEPTED_ROOT / state["snapshot_path"] / "company_facts.csv"
    source_sha, state_sha, workbook_sha = sha(source_csv), sha(state_file), sha(SOURCE)
    frame = pd.read_csv(source_csv, low_memory=False)
    assert len(frame) == 4436 and frame["isin"].is_unique and frame["ticker"].is_unique
    baseline = frame.copy(deep=True)
    for field in FIELD_SOURCES:
        if field not in frame:
            frame[field] = pd.NA
    workbook_mtime = datetime.fromtimestamp(SOURCE.stat().st_mtime, tz=timezone.utc)
    last_source_day = workbook_mtime.date()

    code_index: dict[str, set[int]] = defaultdict(set)
    name_index: dict[str, set[int]] = defaultdict(set)
    for i, row in baseline.iterrows():
        exchange, security, ticker = (str(row[x]).strip().upper() for x in ("primary_exchange", "exchange_security_id", "ticker"))
        if exchange == "NSE":
            code_index[f"NSE:{security}"].add(i)
            code_index[f"NSE:{ticker}"].add(i)
        elif exchange == "BSE":
            code_index[f"BOM:{security}"].add(i)
            code_index[f"BOM:{ticker.removeprefix('BSE-')}"] .add(i)
        for field, prefix in (("screener_nse_code", "NSE:"), ("screener_bse_code", "BOM:")):
            value = row.get(field)
            if not blank(value):
                s = str(value).strip().upper()
                if prefix == "BOM:" and re.fullmatch(r"\d+\.0", s):
                    s = s[:-2]
                code_index[prefix + s].add(i)
        for name in (row["company"], row.get("screener_company_name")):
            normalized = norm_name(name)
            if normalized:
                name_index[normalized].add(i)

    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    source_rows: list[dict] = []
    for rownum, row in enumerate(workbook["CAGR"].iter_rows(values_only=True), 1):
        if rownum <= 2 or not row or not clean_code(row[0]).startswith(("NSE:", "BOM:")):
            continue
        values = {field: valid(field, row[openpyxl.utils.column_index_from_string(col) - 1], last_source_day)
                  for field, (sheet, col) in FIELD_SOURCES.items() if sheet == "CAGR"}
        source_rows.append({"sheet": "CAGR", "row": rownum, "code": clean_code(row[0]), "symbol": None,
                            "name": row[1], "values": values,
                            "cells": {f: col for f, (sheet, col) in FIELD_SOURCES.items() if sheet == "CAGR"}})
    for rownum, row in enumerate(workbook["ATH"].iter_rows(values_only=True), 1):
        if rownum == 1 or not row or blank(row[0]):
            continue
        code = clean_code(row[1]) if len(row) > 1 else ""
        values = {field: valid(field, row[openpyxl.utils.column_index_from_string(col) - 1], last_source_day)
                  for field, (sheet, col) in FIELD_SOURCES.items() if sheet == "ATH"}
        source_rows.append({"sheet": "ATH", "row": rownum, "code": code,
                            "symbol": str(row[0]).strip().upper(), "name": row[2], "values": values,
                            "cells": {f: col for f, (sheet, col) in FIELD_SOURCES.items() if sheet == "ATH"}})
    # I:O is a second, independent keyed return table on the ATH worksheet.
    # Its I codes must never be paired with A:E values by row position.
    for rownum, row in enumerate(workbook["ATH"].iter_rows(min_col=9, max_col=15, values_only=True), 1):
        if rownum == 1 or not row or not clean_code(row[0]).startswith(("NSE:", "BOM:")):
            continue
        values = {field: valid(field, row[openpyxl.utils.column_index_from_string(col) - 9], last_source_day)
                  for field, col in RETURN_FIELDS.items()}
        source_rows.append({"sheet": "ATH_RETURNS", "row": rownum, "code": clean_code(row[0]),
                            "symbol": None, "name": None, "values": values, "cells": RETURN_FIELDS})
    workbook.close()

    matched: dict[tuple[int, str], list[dict]] = defaultdict(list)
    unmatched, ambiguous = [], []
    for item in source_rows:
        candidates = set(code_index.get(item["code"], set())) if item["code"] else set()
        method = "EXACT_EXCHANGE_CODE"
        if not candidates and item["symbol"]:
            candidates = set(code_index.get("NSE:" + item["symbol"], set()))
            method = "EXACT_NSE_SYMBOL"
        if not candidates and item["name"]:
            candidates = set(name_index.get(norm_name(item["name"]), set()))
            method = "UNIQUE_EXACT_NORMALIZED_NAME"
        if item["name"] is None:
            scored = [(i, None) for i in candidates]
            accepted = scored if len(scored) == 1 and method == "EXACT_EXCHANGE_CODE" else []
            method = "EXACT_UNIQUE_EXCHANGE_CODE_NO_NAME"
        else:
            scored = [(i, max(name_score(item["name"], baseline.at[i, "company"]),
                              name_score(item["name"], baseline.at[i, "screener_company_name"]))) for i in candidates]
            accepted = [(i, score) for i, score in scored if score >= (1.0 if method == "UNIQUE_EXACT_NORMALIZED_NAME" else 0.70)]
        descriptor = {"sheet": item["sheet"], "source_row": item["row"], "source_code": item["code"],
                      "source_name": item["name"], "candidate_isins": ";".join(str(baseline.at[i, "isin"]) for i in candidates),
                      "best_name_similarity": max((x[1] for x in scored if x[1] is not None), default=None)}
        if len(accepted) != 1:
            descriptor["reason"] = "NO_IDENTITY_MATCH" if not candidates else "AMBIGUOUS_OR_NAME_MISMATCH"
            (unmatched if not candidates else ambiguous).append(descriptor)
            continue
        i, score = accepted[0]
        item["match_method"], item["name_similarity"], item["matched_index"] = method, score, i
        matched[(i, item["sheet"])].append(item)

    # Duplicate source rows may agree exactly; conflicting values are quarantined.
    fills, conflicts = [], []
    for (i, sheet), items in matched.items():
        fields = list(items[0]["cells"])
        for field in fields:
            observations = [(it["values"][field], it) for it in items if it["values"][field] is not None]
            if not observations:
                continue
            distinct = {str(v) for v, _ in observations}
            if len(distinct) > 1:
                # ATH prices can legitimately differ across NSE/BSE. Retain
                # the row for the accepted identity's primary exchange.
                primary = str(baseline.at[i, "primary_exchange"]).upper()
                prefix = "NSE:" if primary == "NSE" else "BOM:"
                primary_observations = [(v, it) for v, it in observations if it["code"].startswith(prefix)]
                if sheet in {"ATH", "ATH_RETURNS"} and primary_observations and len({str(v) for v, _ in primary_observations}) == 1:
                    observations = primary_observations
                else:
                    conflicts.append({"isin": baseline.at[i, "isin"], "field": field,
                                      "source_rows": ";".join(str(it["row"]) for _, it in observations),
                                      "values": ";".join(sorted(distinct))})
                    continue
            if not blank(frame.at[i, field]):
                # A direct CAGR-sheet value wins over its duplicate in the
                # larger ATH return table. Record any disagreement.
                if sheet == "ATH_RETURNS" and any(x["isin"] == baseline.at[i, "isin"] and x["field"] == field for x in fills):
                    prior = frame.at[i, field]
                    agrees = str(prior) == str(observations[0][0]) if field == "professor_industry" else math.isclose(float(prior), float(observations[0][0]), rel_tol=0, abs_tol=1e-9)
                    if not agrees:
                        conflicts.append({"isin": baseline.at[i, "isin"], "field": field,
                                          "source_rows": str(observations[0][1]["row"]),
                                          "values": f"CAGR:{prior};ATH_RETURNS:{observations[0][0]}"})
                continue
            value, first = observations[0]
            frame.at[i, field] = value
            fills.append({"isin": baseline.at[i, "isin"], "ticker": baseline.at[i, "ticker"],
                          "company": baseline.at[i, "company"], "field": field, "value": value,
                          "provenance": "PROFESSOR_EXCEL", "source_sheet": "ATH" if sheet == "ATH_RETURNS" else sheet,
                          "source_table": sheet,
                          "source_cells": ";".join(f"{it['cells'][field]}{it['row']}" for _, it in observations),
                          "match_method": first["match_method"], "name_similarity": first["name_similarity"],
                          "source_code": first["code"], "workbook_sha256": workbook_sha,
                          "workbook_file_modified_utc": workbook_mtime.isoformat(),
                          "current_only": True})

    # Row-level provenance is a readable index; the fill ledger is authoritative
    # per cell, including exact workbook addresses.
    fill_by_isin: dict[str, list[str]] = defaultdict(list)
    for x in fills:
        fill_by_isin[x["isin"]].append(x["field"])
    frame["professor_excel_filled_fields"] = frame["isin"].map(lambda x: ";".join(sorted(fill_by_isin.get(x, []))))
    frame["professor_excel_source"] = frame["isin"].map(lambda x: "PROFESSOR_EXCEL" if x in fill_by_isin else "")
    frame["professor_excel_source_sha256"] = frame["isin"].map(lambda x: workbook_sha if x in fill_by_isin else "")
    frame["professor_excel_current_only"] = frame["isin"].map(lambda x: bool(x in fill_by_isin))

    # Verify all accepted nonblank cells remain byte-equivalent after CSV roundtrip
    # at the DataFrame level, and only intended fields were changed.
    for col in baseline:
        old = baseline[col]
        protected = ~old.map(blank)
        assert (old[protected].astype(str).to_numpy() == frame.loc[protected, col].astype(str).to_numpy()).all(), col
        if col not in FIELD_SOURCES:
            assert old.equals(frame[col]), col
    before_after = []
    for field in FIELD_SOURCES:
        before = int(baseline[field].map(blank).sum()) if field in baseline else len(baseline)
        after = int(frame[field].map(blank).sum())
        count = sum(x["field"] == field for x in fills)
        assert before - after == count, field
        before_after.append({"field": field, "missing_before": before, "missing_after": after,
                             "filled": count, "nonmissing_before": len(frame)-before, "nonmissing_after": len(frame)-after})
    assert sha(source_csv) == source_sha and sha(state_file) == state_sha and sha(SOURCE) == workbook_sha

    # Copy original CSV cell text, changing only ledger-listed blank cells.
    # This also avoids even sub-ULP float reserialization in accepted fields.
    by_key = {(x["isin"], x["field"]): x["value"] for x in fills}
    new_columns = [x for x in frame.columns if x not in baseline.columns]
    index_by_isin = {isin: i for i, isin in baseline["isin"].items()}
    with source_csv.open("r", encoding="utf-8", newline="") as inp, gzip.open(
        OUT / "current_company_facts_professor_overlay.csv.gz", "wt", encoding="utf-8", newline=""
    ) as output:
        reader = csv.DictReader(inp)
        writer = csv.DictWriter(output, fieldnames=list(reader.fieldnames or []) + new_columns, lineterminator="\n")
        writer.writeheader()
        for row in reader:
            isin = row["isin"]
            for field in FIELD_SOURCES:
                value = by_key.get((isin, field))
                if value is not None:
                    assert field not in row or not row[field].strip(), (isin, field)
                    row[field] = str(value)
            for field in new_columns:
                if field in FIELD_SOURCES:
                    row.setdefault(field, "")
                else:
                    value = frame.at[index_by_isin[isin], field]
                    row[field] = str(value) if not blank(value) else ""
            writer.writerow(row)
    pd.DataFrame(fills).to_csv(OUT / "filled_cells.csv", index=False)
    pd.DataFrame(before_after).to_csv(OUT / "coverage_before_after.csv", index=False)
    pd.DataFrame(unmatched, columns=["sheet", "source_row", "source_code", "source_name", "candidate_isins", "best_name_similarity", "reason"]).to_csv(OUT / "unmatched_rows.csv", index=False)
    pd.DataFrame(ambiguous, columns=["sheet", "source_row", "source_code", "source_name", "candidate_isins", "best_name_similarity", "reason"]).to_csv(OUT / "ambiguous_rows.csv", index=False)
    pd.DataFrame(conflicts, columns=["isin", "field", "source_rows", "values"]).to_csv(OUT / "conflicting_values.csv", index=False)
    summary = {
        "schema": "professor_current_fallback_v1", "source_workbook_sha256": workbook_sha,
        "source_workbook_file_modified_utc": workbook_mtime.isoformat(),
        "accepted_state_sha256": state_sha, "accepted_company_facts_sha256": source_sha,
        "accepted_snapshot": state["snapshot_path"], "company_rows": len(frame),
        "source_rows": dict(Counter(x["sheet"] for x in source_rows)),
        "matched_unique_stocks": len({baseline.at[i, "isin"] for i, _ in matched}),
        "stocks_with_fills": len(fill_by_isin), "filled_cells": len(fills),
        "matched_source_rows": sum(len(v) for v in matched.values()),
        "unmatched_source_rows": len(unmatched), "ambiguous_source_rows": len(ambiguous),
        "conflicting_field_values": len(conflicts), "filled_by_field": dict(Counter(x["field"] for x in fills)),
        "overlay_is_current_only": True, "production_pointer_changed": False,
        "historical_v12_v13_changed": False, "product_baseline_metrics_changed": False,
        "ath_note": "Workbook ATH is a raw price reference; adjusted-history recorded_high and derived baseline days since high are not overwritten.",
        "cagr_note": "Workbook CAGR occupies professor_workbook fields only; product baseline CAGRs remain on their accepted adjusted-history method.",
    }
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    artifact_paths = [OUT / n for n in ("current_company_facts_professor_overlay.csv.gz", "filled_cells.csv", "coverage_before_after.csv",
                                           "unmatched_rows.csv", "ambiguous_rows.csv", "conflicting_values.csv", "summary.json")]
    manifest = {"schema": "professor_current_fallback_manifest_v1", "files": [
        {"path": str(p.relative_to(ROOT)).replace("\\", "/"), "bytes": p.stat().st_size, "sha256": sha(p)} for p in artifact_paths
    ], "source_workbook_sha256": workbook_sha, "accepted_company_facts_sha256": source_sha,
        "build_script_sha256": sha(Path(__file__))}
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
